import os
from openai import OpenAI
from dotenv import load_dotenv
from pypdf import PdfReader
from docx import Document
from datetime import datetime
from pypdf import PdfReader
from pathlib import Path

load_dotenv()
client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

# ---------- FUNCION CHAT ----------
def ask_ai(prompt):
    response = client.chat.completions.create(
        model="gpt-4.1-mini",
        messages=[{"role": "user", "content": prompt}],
        temperature=0.7
    )
    return response.choices[0].message.content


# ---------- COMANDO PDF ----------
def resumir_pdf(ruta):
    reader = PdfReader(ruta)
    texto = ""
    for page in reader.pages:
        texto += page.extract_text()

    prompt = f"Resume este documento de forma clara y profesional:\n\n{texto}"
    return resumir_texto_largo(texto)


# ---------- COMANDO WORD ----------
def resumir_docx(ruta):
    doc = Document(ruta)
    texto = "\n".join([p.text for p in doc.paragraphs])

    prompt = f"Resume este documento de forma clara y profesional:\n\n{texto}"
    return resumir_texto_largo(texto)


# ---------- COMANDO CORREO ----------
def generar_correo(instruccion):
    prompt = f"""
    Escribe un correo profesional basado en esta instrucción:
    {instruccion}
    """
    return ask_ai(prompt)


# ---------- COMANDO MINUTA ----------
def generar_minuta(instruccion):
    prompt = f"""
    Crea una minuta profesional de reunión basada en:
    {instruccion}

    Incluye:
    - Objetivo
    - Puntos tratados
    - Acuerdos
    - Próximos pasos
    """
    return ask_ai(prompt)


# ---------- COMANDO IDEAS ----------
def generar_ideas(instruccion):
    prompt = f"""
    Genera ideas útiles y accionables sobre:
    {instruccion}
    """
    return ask_ai(prompt)


# ---------- LOOP PRINCIPAL ----------
print("🤖 AI Office Assistant PRO listo.")
print("""
Comandos disponibles:
- correo [instrucción]
- minuta [tema]
- ideas [tema]
- pdf archivo.pdf
- docx archivo.docx
- escribir normal para chatear
- salir
""")


# ---------- GUARDAR TXT ----------
def guardar_txt(contenido, prefijo):
    fecha = datetime.now().strftime("%Y-%m-%d_%H-%M")
    nombre = f"{prefijo}_{fecha}.txt"
    with open(nombre, "w") as f:
        f.write(contenido)
    return f"📁 Guardado como {nombre}"


# ---------- GUARDAR WORD ----------
def guardar_docx(contenido, prefijo):
    fecha = datetime.now().strftime("%Y-%m-%d_%H-%M")
    nombre = f"{prefijo}_{fecha}.docx"
    doc = Document()
    doc.add_paragraph(contenido)
    doc.save(nombre)
    return f"📁 Guardado como {nombre}"


# ---------- LECTOR PDF ----------
def leer_pdf(ruta):
    texto = ""
    reader = PdfReader(ruta)
    for page in reader.pages:
        texto += page.extract_text() + "\n"
    return texto


# ---------- RESUMIR TEXTO ----------
def resumir_texto(texto):
    prompt = f"Resume el siguiente documento de forma profesional:\n\n{texto}"

    response = client.chat.completions.create(
        model="gpt-4.1-mini",
        messages=[{"role":"user","content":prompt}]
    )
    return response.choices[0].message.content


# ---------- RESUMIR TEXTO LARGO ----------
def resumir_texto_largo(texto, tamaño_chunk=3000):
    chunks = [texto[i:i+tamaño_chunk] for i in range(0, len(texto), tamaño_chunk)]
    
    resumenes = []
    
    for i, chunk in enumerate(chunks):
        print(f"Procesando parte {i+1}/{len(chunks)}...")
        
        resumen = ask_ai(
            f"Resume este texto de forma clara y profesional:\n\n{chunk}"
        )
        
        resumenes.append(resumen)

    resumen_final = ask_ai(
        "Resume de forma ejecutiva los siguientes resúmenes:\n\n" +
        "\n\n".join(resumenes)
    )

    return resumen_final


# ---------- RESUMIR CARPETA ----------
def resumir_carpeta(ruta_carpeta):
    carpeta = Path(ruta_carpeta)
    
    if not carpeta.exists():
        return "❌ La carpeta no existe."
    
    resumen_total = "RESUMEN GENERAL DE DOCUMENTOS\n\n"
    
    for archivo in carpeta.iterdir():
        try:
            if archivo.suffix.lower() == ".pdf":
                resumen = resumir_pdf(str(archivo))
                resumen_total += f"\n📄 {archivo.name}\n{resumen}\n\n"
            
            elif archivo.suffix.lower() == ".docx":
                resumen = resumir_docx(str(archivo))
                resumen_total += f"\n📝 {archivo.name}\n{resumen}\n\n"
        
        except Exception as e:
            resumen_total += f"\n⚠️ Error leyendo {archivo.name}: {str(e)}\n\n"
    
    return resumen_total


while True:
    user_input = input("Tú: ")

    if user_input.lower() == "salir":
        print("👋 Hasta luego!")
        break

    elif user_input.startswith("pdf "):
        ruta = user_input.replace("pdf ", "")
        print(resumir_pdf(ruta))

    elif user_input.startswith("docx "):
        ruta = user_input.replace("docx ", "")
        print(resumir_docx(ruta))

    elif user_input.startswith("carpeta "):
        ruta = user_input.replace("carpeta ", "")
        resultado = resumir_carpeta(ruta)
        print(resultado)
        print(guardar_docx(resultado, "resumen_carpeta"))

    elif user_input.startswith("correo "):
        instruccion = user_input.replace("correo ", "")
        resultado = generar_correo(instruccion)
        print(resultado)
        print(guardar_docx(resultado, "correo"))
    
    elif user_input.startswith("pdf "):
        archivo = user_input.replace("pdf ", "")
        texto = leer_pdf(archivo)
        resumen = resumir_texto(texto)
        print("\n📄 Resumen del documento:\n")
        print(resumen)

    elif user_input.startswith("minuta "):
        instruccion = user_input.replace("minuta ", "")
        resultado = generar_minuta(instruccion)
        print(resultado)
        print(guardar_docx(resultado, "minuta"))

    elif user_input.startswith("ideas "):
        instruccion = user_input.replace("ideas ", "")
        resultado = generar_ideas(instruccion)
        print(resultado)
        print(guardar_txt(resultado, "ideas"))

    else:
        print(ask_ai(user_input))