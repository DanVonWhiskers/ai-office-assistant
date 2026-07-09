import streamlit as st
from pypdf import PdfReader
import numpy as np
from openai import OpenAI
from dotenv import load_dotenv
from docx import Document
from io import BytesIO
import os
import faiss

# ---------- CONFIG ----------
load_dotenv()

client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

st.set_page_config(
    page_title="AI Office Assistant",
    page_icon="🤖",
    layout="centered"
)

# ---------- FUNCIONES ----------
def ask_ai(prompt):
    response = client.chat.completions.create(
        model="gpt-4.1-mini",
        messages=[
            {
                "role": "system",
                "content": "Eres un asistente profesional experto en resumir documentos."
            },
            {
                "role": "user",
                "content": prompt
            }
        ]
    )

    return response.choices[0].message.content


# ---------- RESUMIR PDF ----------

# ---------- EXTRAER TEXTO PDF ----------
def extraer_texto_pdf(uploaded_file):

    reader = PdfReader(uploaded_file)

    texto = ""

    for page in reader.pages:

        contenido = page.extract_text()

        if contenido:

            texto += contenido

    return texto


# ---------- EXTRAER TEXTO PDF CON PAGINAS ----------
def extraer_paginas_pdf(uploaded_file):

    reader = PdfReader(uploaded_file)

    paginas = []

    for numero, page in enumerate(
        reader.pages,
        start=1
    ):

        contenido = page.extract_text()

        if contenido:

            paginas.append(
                {
                    "pagina": numero,
                    "texto": contenido,
                    "documento": uploaded_file.name
                }
            )

    return paginas


# ---------- RESUMIR TEXTO LARGO ----------
def resumir_texto_largo(texto):

    texto = texto[:12000]

    prompt = f"""
    Resume el siguiente documento de forma clara,
    profesional y estructurada:

    {texto}
    """

    return ask_ai(prompt)

def resumir_pdf(uploaded_file):
    reader = PdfReader(uploaded_file)

    texto = ""

    for page in reader.pages:
        contenido = page.extract_text()

        if contenido:
            texto += contenido

    st.session_state.pdf_text = texto

    texto = texto[:12000]

    prompt = f"""
    Resume el siguiente documento de forma clara, profesional y estructurada:

    {texto}
    """

    return ask_ai(prompt)


# ---------- CREAR DOCX ----------
def crear_docx(resumen):

    doc = Document()

    doc.add_heading("Resumen generado por AI Office Assistant", level=1)

    doc.add_paragraph(resumen)

    buffer = BytesIO()

    doc.save(buffer)

    buffer.seek(0)

    return buffer

# ---------- PREGUNTAR SOBRE PDF ----------
def preguntar_pdf(texto_pdf, pregunta):

    resultado = buscar_chunks_relevantes(
        st.session_state.index,
        st.session_state.chunks,
        pregunta
    )

    contexto = resultado["contexto"]

    paginas = resultado["paginas"]

    historial = ""

    for item in st.session_state.qa_history[-5:]:

        historial += (
            f"Usuario: {item['pregunta']}\n"
            f"Asistente: {item['respuesta']}\n\n"
        )

    prompt = f"""
    Usa el siguiente contexto del documento
    y el historial reciente de conversación.

    Historial:
    {historial}

    Contexto:
    {contexto}

    Pregunta actual:
    {pregunta}

    Si utilizas información del documento,
    indica al final las páginas relevantes.

    Páginas disponibles:
    {paginas}

    Responde de forma clara y profesional.
    """

    return ask_ai(prompt)

# ---------- CREAR EMBEDDINGS ----------
def crear_embedding(texto):

    response = client.embeddings.create(
        model="text-embedding-3-small",
        input=texto
    )

    return response.data[0].embedding
    
# ---------- CREAR ÍNDICE VECTORIAL ----------
def crear_indice_faiss(chunks):

    embeddings = []

    for chunk in chunks:

        embedding = crear_embedding(
        chunk["texto"]
        )

        embeddings.append(embedding)

    embeddings_array = np.array(
        embeddings
    ).astype("float32")

    dimension = embeddings_array.shape[1]

    index = faiss.IndexFlatL2(
        dimension
    )

    index.add(embeddings_array)

    return index, chunks


# ---------- DIVIDIR TEXTO EN CHUNKS ----------
def dividir_texto(texto, tamaño=1000):

    chunks = []

    for i in range(0, len(texto), tamaño):

        chunk = texto[i:i+tamaño]

        chunks.append(chunk)

    return chunks

# ---------- DIVIDIR PAGINAS EN CHUNKS ----------
def dividir_paginas_en_chunks(
    paginas,
    tamaño=1000
):

    chunks = []

    for pagina in paginas:

        texto = pagina["texto"]

        for i in range(
            0,
            len(texto),
            tamaño
        ):

            chunk = texto[i:i+tamaño]

            chunks.append(
                {
                    "texto": chunk,
                    "pagina": pagina["pagina"],
                    "documentos": pagina["documento"]
                }
            )

    return chunks

 # ---------- BUSCAR CHUNKS RELEVANTES ----------
def buscar_chunks_relevantes(
    index,
    chunks,
    pregunta
):

    pregunta_embedding = np.array(
        [crear_embedding(pregunta)]
    ).astype("float32")

    k = 3

    distancias, indices = index.search(
        pregunta_embedding,
        k
    )

    resultados = []
    paginas = set()
    documentos = set()

    for i in indices[0]:

        resultados.append(
            chunks[i]["texto"]
        )

        paginas.add(
            chunks[i]["pagina"]
        )
        
        documentos.add(
            chunks[i]["documentos"]
        )

    return {
        "contexto": "\n\n".join(resultados),
        "paginas": sorted(list(paginas)),
        "documentos": sorted(list(documentos))
    }

# ---------- UI ----------

# ---------- MEMORIA CHAT ----------
if "messages" not in st.session_state:
    st.session_state.messages = []
    
if "qa_history" not in st.session_state:
    st.session_state.qa_history = []

#---------- TÍTULO ----------
st.write("📚 Knowledge Hub")
st.write("Gestiona la información de tus documentos con IA.")

# ---------- SUBIR ARCHIVOS ----------
uploaded_files = st.file_uploader(
    "📄 Arrastra tus PDFs aquí",
    type=["pdf"],
    accept_multiple_files=True
)

if uploaded_files:

    nombres = [
    archivo.name
    for archivo in uploaded_files
    ]

    st.success(
        f"Archivos cargados: {', '.join(nombres)}"
    )

    if st.button("🚀 Analizar documentos"):

        with st.spinner("Procesando documento..."):

            texto_completo = ""

            paginas_completas = []

            for uploaded_file in uploaded_files:

                paginas_pdf = extraer_paginas_pdf(
                    uploaded_file
                )

                paginas_completas.extend(
                    paginas_pdf
                )

                texto_pdf = extraer_texto_pdf(
                    uploaded_file
                )

                texto_completo += texto_pdf + "\n\n"

            resumen = resumir_texto_largo(
                texto_completo
            )

            st.session_state.pdf_text = texto_completo

            chunks = dividir_paginas_en_chunks(
                paginas_completas
            )

            index, chunks = crear_indice_faiss(
                chunks
            )

            st.session_state.index = index
            st.session_state.chunks = chunks

            st.session_state.resumen = resumen

            st.session_state.messages.append(
                {
                    "role": "user",
                    "content": (
                        f"Subí archivos: "
                        f"{', '.join(nombres)}"
                    )
                }
            )

            st.session_state.messages.append(
                {
                    "role": "assistant",
                    "content": resumen
                }
            )

            st.rerun()

if "resumen" in st.session_state:

    docx_file = crear_docx(
        st.session_state.resumen
    )

    st.download_button(
        label="📥 Descargar resumen en Word",
        data=docx_file,
        file_name="resumen.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

# ---------- MOSTRAR CHAT ----------
for message in st.session_state.messages:

    with st.chat_message(message["role"]):

        st.markdown(message["content"])

# ---------- ACCIONES ----------

if (
    st.session_state.messages
    or st.session_state.qa_history
):

    col1, col2 = st.columns([5,1])

    with col2:

        if st.button("🗑️ Nueva conversación"):

            st.session_state.qa_history = []

            st.rerun()

# --------- MENSAJE BIENVENIDA ----------

if "resumen" not in st.session_state:

    st.info(

        """

📚 Bienvenido a Knowledge Hub

Transforma tus documentos en una fuente de conocimiento consultable.

Ejemplos:

• Resume este documento

• ¿Cuáles son los requisitos?

• Enumera las tecnologías mencionadas

• ¿Qué responsabilidades tiene el puesto?

• ¿Qué riesgos identifica el documento?

        """

    )

# ---------- HISTORIAL ----------
for item in st.session_state.qa_history:

    with st.chat_message("user"):
        st.markdown(item["pregunta"])

    with st.chat_message("assistant"):
        st.markdown(item["respuesta"])


# ---------- CHAT ----------
st.divider()

pregunta = st.chat_input(
    "Pregunta sobre tus documentos..."
)

if pregunta:

    with st.spinner("Pensando..."):

        respuesta = preguntar_pdf(
            st.session_state.pdf_text,
            pregunta
        )

    st.session_state.qa_history.append(
        {
            "pregunta": pregunta,
            "respuesta": respuesta
        }
    )

    st.rerun()